import os
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from engines.escalation import XLSX_COLUMNS, prefill_demand_notice


@pytest.fixture
def inputs(tmp_path):
    out_dir = tmp_path / "out"
    out_dir.mkdir()
    invoice = {
        "invoice_id": "INV-2026-0042",
        "amount": 500_000.0,
        "due_date": date(2026, 5, 15),
        "buyer_name": "Acme Buyer Pvt Ltd",
        "appointed_day": date(2026, 5, 16),
    }
    supplier = {
        "name": "Widgets & Co",
        "udyam_number": "UDYAM-MH-12-0001234",
        "flow": "full",
    }
    interest_result = {
        "principal": 500_000.0,
        "interest": 31_820.02,
        "total_due": 531_820.02,
        "statutory_rate_pa": 0.165,
        "bank_rate_used": 0.055,
        "schedule": [],
    }
    return str(out_dir), invoice, supplier, interest_result


def test_returns_pending_human_approval(inputs):
    out_dir, invoice, supplier, ir = inputs
    r = prefill_demand_notice(invoice, supplier, ir, output_dir=out_dir)
    assert r["status"] == "PENDING_HUMAN_APPROVAL"
    assert r["invoice_id"] == "INV-2026-0042"


def test_pdf_file_created(inputs):
    out_dir, invoice, supplier, ir = inputs
    r = prefill_demand_notice(invoice, supplier, ir, output_dir=out_dir)
    assert os.path.isfile(r["pdf_path"])


def test_xlsx_file_created(inputs):
    out_dir, invoice, supplier, ir = inputs
    r = prefill_demand_notice(invoice, supplier, ir, output_dir=out_dir)
    assert os.path.isfile(r["xlsx_path"])


def test_pdf_contains_human_review_marker(inputs):
    out_dir, invoice, supplier, ir = inputs
    r = prefill_demand_notice(invoice, supplier, ir, output_dir=out_dir)
    raw = Path(r["pdf_path"]).read_bytes().decode("latin-1")
    assert "REQUIRES HUMAN REVIEW" in raw


def test_pdf_contains_invoice_id(inputs):
    out_dir, invoice, supplier, ir = inputs
    r = prefill_demand_notice(invoice, supplier, ir, output_dir=out_dir)
    raw = Path(r["pdf_path"]).read_bytes().decode("latin-1")
    assert "INV-2026-0042" in raw


def test_xlsx_headers_in_order(inputs):
    out_dir, invoice, supplier, ir = inputs
    r = prefill_demand_notice(invoice, supplier, ir, output_dir=out_dir)
    wb = load_workbook(r["xlsx_path"])
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    assert headers == XLSX_COLUMNS


def test_xlsx_data_row_values(inputs):
    out_dir, invoice, supplier, ir = inputs
    r = prefill_demand_notice(invoice, supplier, ir, output_dir=out_dir)
    wb = load_workbook(r["xlsx_path"])
    ws = wb.active
    row = {h: v for h, v in zip([c.value for c in ws[1]], [c.value for c in ws[2]])}
    assert row["principal_inr"] == 500_000.0
    assert row["interest_inr"] == 31_820.02
    assert row["total_due_inr"] == 531_820.02
    assert row["udyam_number"] == "UDYAM-MH-12-0001234"


def test_reminders_only_supplier_raises(inputs):
    out_dir, invoice, supplier, ir = inputs
    supplier["flow"] = "reminders_only"
    with pytest.raises(ValueError):
        prefill_demand_notice(invoice, supplier, ir, output_dir=out_dir)


def test_missing_output_dir_raises(inputs):
    _, invoice, supplier, ir = inputs
    with pytest.raises(ValueError):
        prefill_demand_notice(invoice, supplier, ir, output_dir="no/such/dir")


def test_calling_twice_overwrites(inputs):
    out_dir, invoice, supplier, ir = inputs
    r1 = prefill_demand_notice(invoice, supplier, ir, output_dir=out_dir)
    r2 = prefill_demand_notice(invoice, supplier, ir, output_dir=out_dir)
    assert r1["pdf_path"] == r2["pdf_path"]
    assert os.path.isfile(r2["pdf_path"])
    assert os.path.isfile(r2["xlsx_path"])
